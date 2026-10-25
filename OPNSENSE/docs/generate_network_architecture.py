#!/usr/bin/env python3
"""Generate Peak Energy network architecture PDF + Word documents."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUT = Path(__file__).resolve().parent
DIAGRAM_PNG = OUT / "Peak-Energy-Network-Architecture.png"
PDF_PATH = OUT / "Peak-Energy-Network-Architecture.pdf"
DOCX_PATH = OUT / "Peak-Energy-Network-Architecture.docx"

TODAY = date.today().isoformat()

VLANS = [
    ["VLAN / Role", "Subnet", "Gateway (FW)", "DHCP pool", "Interface"],
    ["MGMT (admin gear)", "10.80.99.0/24", "10.80.99.1", "10.80.99.101–200", "lan / re0"],
    ["WIRED_CORP", "10.80.100.0/24", "10.80.100.1", "10.80.100.101–200", "opt2 / VLAN 100"],
    ["WL_CORP (PETCPL)", "10.80.101.0/24", "10.80.101.1", "10.80.101.101–200", "opt3 / VLAN 101"],
    ["WL_GUEST", "10.80.102.0/24", "10.80.102.1", "10.80.102.101–200", "opt4 / VLAN 102"],
    ["WireGuard VPN", "10.80.200.0/28", "10.80.200.1", "static peers", "opt6 / wg0"],
]

WAN = [
    ["WAN", "Public IP", "Gateway name", "Notes"],
    ["Airtel Fiber", "125.19.224.18/30", "WAN_AIRTEL", "Primary; WG endpoint default"],
    ["Jio RF", "47.247.169.94/30", "WAN_JIO", "Failover / alternate WG endpoint"],
    ["WAN_LB", "—", "WAN_LB", "Dual-WAN load balance + failover"],
]

MGMT_HOSTS = [
    ["IP", "Hostname", "Role"],
    ["10.80.99.1", "PEAK-CORP-FW / OpnSense", "Firewall GUI :4444, SSH, Unbound DNS"],
    ["10.80.99.2", "sg2428p", "Omada switch SG2428P"],
    ["10.80.99.3", "oc200", "Omada OC200 controller https://10.80.99.3/"],
    ["10.80.99.21", "eap670-1", "Wi-Fi AP EAP670"],
    ["10.80.99.22", "eap670-2", "Wi-Fi AP EAP670"],
    ["10.80.99.23", "eap670-3", "Wi-Fi AP EAP670"],
    ["10.80.99.24", "eap225-3", "Wi-Fi AP EAP225"],
]

WIRED_HOSTS = [
    ["IP", "Hostname", "Role"],
    ["10.80.100.1", "—", "WIRED_CORP SVI / DNS"],
    ["10.80.100.10", "biometric-server", "ACS / Attendance SQL"],
    ["10.80.100.11", "biometric-entry", "Biometric entry terminal"],
    ["10.80.100.12", "biometric-exit", "Biometric exit terminal"],
    ["10.80.100.21", "laser-printer", "HP LaserJet (Ethernet)"],
    ["10.80.100.51", "peak-ws-1", "Workstation 1"],
    ["10.80.100.52", "peak-ws-2", "Workstation 2"],
    ["10.80.100.53", "peak-ws-3", "Workstation 3"],
    ["10.80.100.54", "dev-ws-1", "PeakPulse / peakpulse-dev.peakenergy.asia"],
    ["10.80.100.55", "peak801-in-conf", "Conference room PC"],
]

WIFI_HOSTS = [
    ["IP", "Hostname", "Role"],
    ["10.80.101.1", "—", "WL_CORP SVI / DNS"],
    ["10.80.101.21", "laser-printer", "HP LaserJet (Wi-Fi)"],
    ["10.80.101.22", "printer-md", "MD HP Smart Tank Wi-Fi"],
    ["10.80.101.30", "tv-conferance", "TV Conference Room"],
    ["10.80.101.31", "tv-kalaam", "TV Abdul Kalaam"],
    ["10.80.101.32", "tv-ratan", "TV Ratan Tata"],
    ["10.80.101.33", "tv-onm", "TV Operations & Mgmt"],
    ["10.80.101.34", "tv-md", "TV MD"],
]

WG_PEERS = [
    ["Peer", "Tunnel IP", "Access"],
    ["PEAK-WG server", "10.80.200.1/28", "OPNsense wg0 listen UDP 51820"],
    ["binny.andrews", "10.80.200.2/32", "Admin: FW GUI/SSH + MGMT + CORP"],
    ["admin", "10.80.200.3/32", "Admin: FW GUI/SSH + MGMT + CORP"],
    ["jagadeshwar", "10.80.200.4/32", "AMC: CORP .100/.101; RDP allowed"],
    ["venu.gopal.reddy", "10.80.200.5/32", "AMC: CORP only; RDP blocked"],
    ["poovarasu", "10.80.200.6/32", "AMC: CORP only; RDP blocked"],
    ["sanoj.james", "10.80.200.7/32", "AMC: CORP only; RDP blocked"],
    ["sampada", "10.80.200.8/32", "AMC: CORP only; RDP blocked"],
    ["sharan", "10.80.200.9/32", "AMC: CORP only; RDP blocked"],
]

POLICY = [
    ["Topic", "Current policy"],
    ["PeakPulse public WAN", "DISABLED — HTTP/HTTPS DNAT kept but turned off on both WANs"],
    ["PeakPulse access", "On-site LAN or WireGuard only → 10.80.100.54 (dev-ws-1)"],
    ["WG DNS", "New peers use 10.80.200.1 (Unbound on WIREGUARD). Legacy may use 10.80.100.1. Both are allowed."],
    ["WG AllowedIPs", "10.80.0.0/16 split tunnel (corp only, not full internet)"],
    ["WG endpoints", "125.19.224.18:51820 (Airtel, default) or 47.247.169.94:51820 (Jio)"],
    ["WG interface", "opt6 / wg0, description WIREGUARD. Routes disabled; firewall does the policy."],
    ["MGMT via VPN", "Only WG_FW_ADMINS (binny.andrews .2 and admin .3)"],
    ["FW GUI / SSH via VPN", "Only WG_FW_ADMINS. GUI listens on MGMT, WIRED, WL_CORP, and WIREGUARD — never WAN."],
    ["AMC via VPN", "CORP only: 10.80.100.0/24 and 10.80.101.0/24. No MGMT, no firewall GUI."],
    ["RDP over VPN", "TCP 3389 blocked for Venu, Poovarasu, Sanoj, Sampada, Sharan. Allowed for binny, admin, Jagadeshwar."],
    ["ACS SQL from Internet", "KEKA_HR only (cin3.a.keka.com) → 10.80.100.10:1433 on both WANs"],
    ["Akrivia", "AKRIVIA_HR alias kept, disabled"],
    ["Corp Wi-Fi MAC list", "WL_CORP_MACS rules exist but are DISABLED. PETCPL is open to the VLAN."],
    ["Guest Wi-Fi", "SSID PETCPL-GUEST, VLAN 102, captive portal. Internet Mon–Fri 07:00–19:00. No internal access."],
    ["On-site vs remote", "On PETCPL or wired corp, Peak Energy VPN OFF. Remote, VPN ON."],
]

DNS_HOSTS = [
    ["Name", "Address", "Use"],
    ["fw.peakenergy.asia", "10.80.99.1", "Firewall (MGMT)"],
    ["peak-corp-fw.peakenergy.asia", "10.80.99.1", "Firewall (MGMT)"],
    ["fw-corp.peakenergy.asia", "10.80.100.1", "Firewall (WIRED_CORP)"],
    ["acs.peakenergy.asia", "10.80.100.10", "Attendance / ACS SQL"],
    ["oc200.peakenergy.asia", "10.80.99.3", "Omada controller"],
    ["app.peakenergy.asia", "10.80.100.54", "PeakPulse"],
    ["peakpulse-dev.peakenergy.asia", "10.80.100.54", "PeakPulse"],
]

BIO = [
    ["Host", "Address", "Role today"],
    ["biometric-server", "10.80.100.10", "Active ACS. Alias ACS_SERVER. Keka SQL target."],
    ["biometric-entry", "10.80.100.11", "Entry terminal. Same VLAN as ACS."],
    ["biometric-exit", "10.80.100.12", "Exit terminal. Same VLAN as ACS."],
]


def draw_diagram(path: Path) -> None:
    """Draw architecture PNG with Pillow."""
    from PIL import Image as PILImage, ImageDraw, ImageFont

    W, H = 1600, 1080
    img = PILImage.new("RGB", (W, H), "#F7F8FA")
    d = ImageDraw.Draw(img)

    def font(size: int):
        for name in (
            "C:/Windows/Fonts/segoeui.ttf",
            "C:/Windows/Fonts/arial.ttf",
            "C:/Windows/Fonts/calibri.ttf",
        ):
            try:
                return ImageFont.truetype(name, size)
            except OSError:
                continue
        return ImageFont.load_default()

    f_title = font(28)
    f_h = font(16)
    f_b = font(13)
    f_s = font(11)

    def box(xy, fill, outline, title, lines, title_fill="#111"):
        x1, y1, x2, y2 = xy
        d.rounded_rectangle(xy, radius=12, fill=fill, outline=outline, width=2)
        d.text((x1 + 14, y1 + 10), title, fill=title_fill, font=f_h)
        yy = y1 + 36
        for line in lines:
            d.text((x1 + 14, yy), line, fill="#222", font=f_s)
            yy += 18

    def arrow(a, b, color="#555"):
        d.line([a, b], fill=color, width=3)

    d.text((40, 24), "Peak Energy — Corporate Network Architecture", fill="#0B3D5C", font=f_title)
    d.text((40, 60), f"PEAK-CORP-FW (OPNsense)  ·  Document date {TODAY}", fill="#555", font=f_b)

    # Internet
    box((60, 100, 280, 200), "#E8F1FF", "#2F6FED", "Internet / ISP", [
        "Clients / Keka cloud",
        "Remote VPN users",
    ])

    # WANs
    box((360, 90, 620, 170), "#FFF4E5", "#D97706", "WAN_AIRTEL", [
        "125.19.224.18/30",
        "WG :51820 (primary)",
    ])
    box((360, 190, 620, 270), "#FFF4E5", "#D97706", "WAN_JIO", [
        "47.247.169.94/30",
        "WG :51820 (alternate)",
    ])

    arrow((280, 150), (360, 130))
    arrow((280, 160), (360, 230))

    # Firewall
    box((680, 100, 1040, 280), "#E8F8EF", "#059669", "PEAK-CORP-FW (OPNsense)", [
        "GUI https://10.80.99.1:4444",
        "Unbound DNS · Kea DHCP",
        "WAN_LB dual-WAN",
        "WireGuard PEAK-WG 10.80.200.1/28",
        "Public PeakPulse NAT: DISABLED",
    ])
    arrow((620, 150), (680, 170))
    arrow((620, 230), (680, 200))

    # Switch
    box((1100, 100, 1540, 270), "#F3E8FF", "#7C3AED", "Omada SG2428P  10.80.99.2", [
        "L2 switch · VLANs 99/100/101/102",
        "OC200  10.80.99.3",
        "APs EAP670 x3 + EAP225",
        "SSIDs PETCPL / PETCPL-GUEST",
    ])
    arrow((1040, 190), (1100, 190))

    # MGMT
    box((60, 320, 400, 520), "#FEE2E2", "#DC2626", "MGMT  10.80.99.0/24", [
        ".1  Firewall",
        ".2  Switch SG2428P",
        ".3  OC200",
        ".21–.23  EAP670 APs",
        ".24  EAP225",
        "Admin gear only",
        "WG: binny/admin only",
    ])

    # WIRED
    box((430, 320, 800, 540), "#DBEAFE", "#2563EB", "WIRED_CORP  10.80.100.0/24", [
        ".1  SVI / DNS",
        ".10 ACS (active)",
        ".11 entry  .12 exit",
        ".21 laser printer",
        ".51–.53 workstations",
        ".54 PeakPulse (dev-ws-1)",
        ".55 conference PC",
        "DHCP .101–.200",
    ])

    # WL CORP
    box((830, 320, 1180, 520), "#D1FAE5", "#059669", "WL_CORP (PETCPL)  10.80.101.0/24", [
        ".1  SVI / DNS",
        "SSID PETCPL",
        "TVs .30–.34",
        "Printers .21 / .22",
        "DHCP .101–.200",
        "MAC list kept, disabled",
    ])

    # GUEST
    box((1210, 320, 1540, 520), "#FEF3C7", "#D97706", "WL_GUEST  10.80.102.0/24", [
        "SSID PETCPL-GUEST",
        "Captive portal",
        "Internet Mon–Fri 07–19",
        "No internal access",
        "DHCP .101–.200",
    ])

    arrow((1290, 270), (1290, 320))

    # VPN
    box((40, 560, 780, 820), "#E0E7FF", "#4F46E5", "Peak Energy VPN (WireGuard PEAK-WG)", [
        "UDP 51820 both WANs · opt6 / wg0 · 10.80.200.0/28",
        "Split tunnel AllowedIPs = 10.80.0.0/16",
        "DNS 10.80.200.1 (Unbound on the tunnel)",
        "Admins .2/.3: firewall GUI/SSH + MGMT + CORP",
        "AMC .4–.9: CORP .100/.101 only (PeakPulse)",
        "RDP: .4 allowed; .5–.9 blocked",
        "Client: Peak Energy VPN.exe · On-site: OFF",
    ])
    arrow((860, 270), (400, 560), "#4F46E5")

    box((810, 560, 1560, 780), "#ECFDF5", "#047857", "Biometrics / Keka", [
        "ACS .10 (active)  ·  entry .11  ·  exit .12",
        "Keka (cin3.a.keka.com) → WAN :1433 → .10",
        "AKRIVIA_HR disabled",
        "acs.peakenergy.asia → 10.80.100.10",
        "Punches stay on VLAN 100 (switch, not firewall)",
    ])

    box((40, 850, 1560, 980), "#F1F5F9", "#64748B", "Names and apps (internal DNS only)", [
        "peakpulse-dev / app .peakenergy.asia → 10.80.100.54   ·   public HTTP/HTTPS DNAT disabled",
        "fw / peak-corp-fw → 10.80.99.1    fw-corp → 10.80.100.1    oc200 → 10.80.99.3",
        "Domain peakenergy.asia · Unbound on MGMT, WIRED_CORP, WL_CORP, and WIREGUARD",
    ])

    d.text((40, 1020), "Source: live PEAK-CORP-FW policy + WireGuard peer pack · Peak Energy IT", fill="#888", font=f_s)
    img.save(path, "PNG", optimize=True)


def style_table(t: Table, header=True):
    style = TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B3D5C")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ]
    )
    t.setStyle(style)


def flow_table(rows: list[list[str]], col_widths: list[float]) -> Table:
    cell = ParagraphStyle(
        "Cell",
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#111111"),
    )
    head = ParagraphStyle(
        "HeadCell",
        parent=cell,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )
    data = []
    for i, row in enumerate(rows):
        style = head if i == 0 else cell
        data.append([Paragraph(str(val).replace("&", "&amp;"), style) for val in row])
    table = Table(data, colWidths=col_widths)
    style_table(table)
    return table


def build_pdf(diagram: Path, path: Path) -> None:
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Peak Energy Network Architecture",
        author="Peak Energy IT",
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=16, textColor=colors.HexColor("#0B3D5C"), spaceAfter=8)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12, textColor=colors.HexColor("#0B3D5C"), spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=9, leading=12)
    story = []

    story.append(Paragraph("Peak Energy — Corporate Network Architecture", h1))
    story.append(Paragraph(f"Firewall: <b>PEAK-CORP-FW</b> (OPNsense) · Generated {TODAY}", body))
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        "This document describes the live network: VLANs, WAN links, Omada Wi‑Fi, "
        "biometrics and Keka, WireGuard (Peak Energy VPN), and access policy for PeakPulse.",
        body,
    ))
    story.append(Spacer(1, 8))

    # Landscape-ish diagram scaled to page width
    iw, ih = 170 * mm, 114.75 * mm
    story.append(Image(str(diagram), width=iw, height=ih))
    story.append(PageBreak())

    story.append(Paragraph("1. WAN links", h2))
    story.append(flow_table(WAN, [100, 120, 100, 180]))

    story.append(Paragraph("2. VLANs / subnets", h2))
    story.append(flow_table(VLANS, [110, 90, 80, 100, 110]))

    story.append(Paragraph("3. MGMT hosts (10.80.99.0/24)", h2))
    story.append(flow_table(MGMT_HOSTS, [90, 140, 270]))

    story.append(Paragraph("4. WIRED_CORP hosts (10.80.100.0/24)", h2))
    story.append(flow_table(WIRED_HOSTS, [90, 140, 270]))

    story.append(Paragraph("5. WL_CORP hosts (10.80.101.0/24)", h2))
    story.append(flow_table(WIFI_HOSTS, [90, 140, 270]))

    story.append(PageBreak())
    story.append(Paragraph("6. WireGuard / Peak Energy VPN", h2))
    story.append(Paragraph(
        "The Windows client is "
        "<b>Peak Energy VPN</b>. On PETCPL or wired corp the tunnel must be off, "
        "because AllowedIPs sends 10.80.0.0/16 into WireGuard.",
        body,
    ))
    story.append(flow_table(WG_PEERS, [120, 100, 280]))

    story.append(Paragraph("7. Biometrics and Keka", h2))
    story.append(Paragraph(
        "Attendance stays on WIRED_CORP. Entry and exit terminals talk to the ACS "
        "on the same VLAN, so that traffic does not pass the firewall. "
        "Keka cloud pulls SQL from the public WAN address on TCP 1433. "
        "NAT on Airtel and Jio forwards that only from KEKA_HR to 10.80.100.10, "
        "the active ACS. The old Akrivia source alias is kept and disabled.",
        body,
    ))
    story.append(flow_table(BIO, [130, 90, 280]))

    story.append(Paragraph("8. Internal DNS (Unbound)", h2))
    story.append(flow_table(DNS_HOSTS, [180, 90, 230]))

    story.append(Paragraph("9. Security & access policy", h2))
    story.append(flow_table(POLICY, [140, 360]))

    story.append(Paragraph("10. Logical traffic paths", h2))
    story.append(Paragraph(
        "<b>On-site (PETCPL / wired):</b> Clients use the VLAN gateway. "
        "Peak Energy VPN stays <b>OFF</b> so 10.80.x is not forced into the tunnel.<br/>"
        "<b>Remote:</b> Peak Energy VPN <b>ON</b> → UDP 51820 → PEAK-WG. "
        "AMC peers reach CORP (.100/.101), including PeakPulse and the ACS. "
        "Admins also reach MGMT (.99) and the firewall GUI. "
        "RDP is blocked for every AMC peer except Jagadeshwar.<br/>"
        "<b>Keka:</b> cin3.a.keka.com → WAN TCP 1433 → 10.80.100.10.<br/>"
        "<b>Guest:</b> PETCPL-GUEST on VLAN 102, captive portal, blocked from MGMT and CORP, "
        "internet only on the weekday schedule.",
        body,
    ))

    story.append(Spacer(1, 12))
    story.append(Paragraph(
        "<i>Confidential — Peak Energy IT. Live PEAK-CORP-FW policy and WireGuard peer pack.</i>",
        body,
    ))
    doc.build(story)


def set_cell_shading(cell, hex_color: str):
    from docx.oxml import OxmlElement

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")


def add_docx_table(doc: Document, rows: list[list[str]]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i == 0:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            if i == 0:
                set_cell_shading(cell, "0B3D5C")
    doc.add_paragraph()


def build_docx(diagram: Path, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = doc.add_heading("Peak Energy — Corporate Network Architecture", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    run = p.add_run(f"Firewall: PEAK-CORP-FW (OPNsense)  ·  Generated {TODAY}")
    run.font.size = Pt(11)

    doc.add_paragraph(
        "This document describes the live network: VLANs, WAN links, Omada Wi‑Fi, "
        "biometrics and Keka, WireGuard (Peak Energy VPN), and access policy for PeakPulse."
    )

    doc.add_heading("Architecture diagram", level=1)
    doc.add_picture(str(diagram), width=Inches(7.0))

    doc.add_heading("1. WAN links", level=1)
    add_docx_table(doc, WAN)

    doc.add_heading("2. VLANs / subnets", level=1)
    add_docx_table(doc, VLANS)

    doc.add_heading("3. MGMT hosts (10.80.99.0/24)", level=1)
    add_docx_table(doc, MGMT_HOSTS)

    doc.add_heading("4. WIRED_CORP hosts (10.80.100.0/24)", level=1)
    add_docx_table(doc, WIRED_HOSTS)

    doc.add_heading("5. WL_CORP hosts (10.80.101.0/24)", level=1)
    add_docx_table(doc, WIFI_HOSTS)

    doc.add_heading("6. WireGuard / Peak Energy VPN", level=1)
    doc.add_paragraph(
        "The Windows client is Peak Energy VPN. "
        "On PETCPL or wired corp the tunnel must be off, because AllowedIPs sends "
        "10.80.0.0/16 into WireGuard."
    )
    add_docx_table(doc, WG_PEERS)

    doc.add_heading("7. Biometrics and Keka", level=1)
    doc.add_paragraph(
        "Attendance stays on WIRED_CORP. Entry and exit terminals talk to the ACS "
        "on the same VLAN, so that traffic does not pass the firewall. "
        "Keka cloud pulls SQL from the public WAN address on TCP 1433. "
        "NAT on Airtel and Jio forwards that only from KEKA_HR to 10.80.100.10, "
        "the active ACS. The old Akrivia source alias is kept and disabled."
    )
    add_docx_table(doc, BIO)

    doc.add_heading("8. Internal DNS (Unbound)", level=1)
    add_docx_table(doc, DNS_HOSTS)

    doc.add_heading("9. Security and access policy", level=1)
    add_docx_table(doc, POLICY)

    doc.add_heading("10. Logical traffic paths", level=1)
    doc.add_paragraph(
        "On-site (PETCPL / wired): clients use the VLAN gateway. "
        "Peak Energy VPN stays OFF so 10.80.x is not forced into the tunnel."
    )
    doc.add_paragraph(
        "Remote: Peak Energy VPN ON, then UDP 51820 to PEAK-WG. "
        "AMC peers reach CORP (.100 and .101), including PeakPulse and the ACS. "
        "Admins also reach MGMT (.99) and the firewall GUI. "
        "RDP is blocked for every AMC peer except Jagadeshwar."
    )
    doc.add_paragraph(
        "Keka: cin3.a.keka.com to WAN TCP 1433, then 10.80.100.10."
    )
    doc.add_paragraph(
        "Guest: PETCPL-GUEST on VLAN 102, captive portal, blocked from MGMT and CORP, "
        "internet only on the weekday schedule."
    )

    doc.add_paragraph()
    note = doc.add_paragraph()
    r = note.add_run(
        "Confidential — Peak Energy IT. Live PEAK-CORP-FW policy and WireGuard peer pack. Generated " + TODAY + "."
    )
    r.italic = True
    r.font.size = Pt(9)

    doc.save(str(path))


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    print("Drawing diagram…")
    draw_diagram(DIAGRAM_PNG)
    print("Writing PDF…")
    build_pdf(DIAGRAM_PNG, PDF_PATH)
    print("Writing Word…")
    build_docx(DIAGRAM_PNG, DOCX_PATH)
    print(f"PNG:  {DIAGRAM_PNG}")
    print(f"PDF:  {PDF_PATH}")
    print(f"DOCX: {DOCX_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
